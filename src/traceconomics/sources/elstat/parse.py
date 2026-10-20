"""Read ELSTAT's CPI table 07: one year's index for every item, with that year's weights.

One file per year (2010 onward, re-issued in COICOP 2018 with 2020=100),
in English and in Greek. `.xls` up to 2025, `.xlsx` from 2026. The layout:

    row 1   title, naming the year ("... for 2025 (Base year: 2020=100,0)")
    row 3   headers: COICOP18 code | label | Weights (‰) | DEC-2024 | JAN ... DEC | ANNUAL AVERAGE
    then    one row per item, grouped by division with blank and heading rows between,
            the divisions repeated in a summary block at the end, and the overall
            index on a row with no code.

The December column (the link month) is there from 2014 only; for 2010-2013
the previous December is in the previous year's file, on the same base.
Weights are per mille of the whole basket at every level; indices and weights
are kept at the full precision of the file.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import xlrd

from traceconomics.checks import FormatError, examples, require

# Bump when parsing changes in a way that changes the output: the store then re-parses.
PARSER_VERSION = "1"

SHEET_NAMES = ("INDICES", "ΔΕΙΚΤΕΣ")
MONTHS = {
    "JAN": 1, "FEB": 2, "MAR": 3, "APR": 4, "MAY": 5, "MAI": 5, "JUN": 6, "JUL": 7, "AUG": 8, "SEP": 9,
    "OCT": 10, "NOV": 11, "DEC": 12,
    "ΙΑΝ": 1, "ΦΕΒ": 2, "ΜΑΡ": 3, "ΑΠΡ": 4, "ΜΑΪ": 5, "ΜΑΙ": 5, "ΙΟΥΝ": 6, "ΙΟΥΛ": 7, "ΑΥΓ": 8,
    "ΣΕΠ": 9, "ΟΚΤ": 10, "ΝΟΕ": 11, "ΔΕΚ": 12,
}
LINK_MONTH = re.compile(r"^(DEC|ΔΕΚ)-(\d{4})$")
YEAR_IN_TITLE = re.compile(r"(?:for|έτους)\s+(\d{4})")
OVERALL_LABELS = ("OVERALL INDEX", "ΓΕΝΙΚΟΣ ΔΕΙΚΤΗΣ")
CODE = re.compile(r"^\d{2,6}$")
TOTAL = "TOTAL"


@dataclass
class Table07:
    year: int
    language: str                                   # "en" or "el"
    weights: dict[str, float] = field(default_factory=dict)       # code -> per mille
    weight_text: dict[str, str] = field(default_factory=dict)
    indices: dict[str, dict[str, float]] = field(default_factory=dict)   # code -> {"YYYY-MM": index}
    annual_average: dict[str, float] = field(default_factory=dict)
    labels: dict[str, str] = field(default_factory=dict)

    def rows(self):
        """Long form: (freq, unit, code, ref_period, value). Unit I20 = index 2020=100; PM = per mille."""
        for code, w in self.weights.items():
            yield "A", "PM", code, str(self.year), w
        for code, series in self.indices.items():
            for period, v in series.items():
                yield "M", "I20", code, period, v
        for code, v in self.annual_average.items():
            yield "A", "I20", code, str(self.year), v


def _read_rows(path: Path) -> list[list]:
    if path.suffix.lower() == ".xls":
        book = xlrd.open_workbook(path, on_demand=True)
        names = [n for n in book.sheet_names() if n.strip() in SHEET_NAMES]
        require(bool(names), f"{path.name}: no sheet named {' or '.join(SHEET_NAMES)}: {book.sheet_names()}")
        sheet = book.sheet_by_name(names[0])
        return [sheet.row_values(i) for i in range(sheet.nrows)]
    if path.suffix.lower() == ".xlsx":
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            names = [n for n in book.sheetnames if n.strip() in SHEET_NAMES]
            require(bool(names), f"{path.name}: no sheet named {' or '.join(SHEET_NAMES)}: {book.sheetnames}")
            return [list(r) for r in book[names[0]].iter_rows(values_only=True)]
        finally:
            book.close()
    raise FormatError(f"{path.name}: not an Excel file")


def _text(v) -> str:
    return "" if v is None else str(v).strip()


def _number(v) -> float | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    raise FormatError(f"expected a number, got {v!r}")


def parse(path: Path) -> Table07:
    rows = _read_rows(path)
    require(len(rows) > 10, f"{path.name}: only {len(rows)} rows")
    title = " ".join(_text(v) for v in rows[0])
    m = YEAR_IN_TITLE.search(title)
    require(m is not None, f"{path.name}: no year in the title: {title[:100]!r}")
    year = int(m.group(1))
    require("2020=100" in title, f"{path.name}: title doesn't say 2020=100 (base changed?): {title[:100]!r}")
    language = "el" if "έτους" in title else "en"

    head_at = next((i for i, r in enumerate(rows[:8]) if _text(r[0]).upper().startswith("COICOP18")), None)
    require(head_at is not None, f"{path.name}: no COICOP18 header row")
    header = [_text(v).upper() for v in rows[head_at]]
    require("‰" in header[2], f"{path.name}: third column isn't the weights (‰): {header[2]!r}")

    cols: dict[int, str] = {}   # column -> period
    average_col = None
    for j, h in enumerate(header[3:], start=3):
        if not h:
            continue
        if lm := LINK_MONTH.match(h):
            require(int(lm.group(2)) == year - 1, f"{path.name}: link month {h} in the {year} table")
            cols[j] = f"{year - 1}-12"
        elif h in MONTHS:
            cols[j] = f"{year}-{MONTHS[h]:02d}"
        elif "ANNUAL" in h or "ΕΤΗΣΙΟΣ" in h:
            average_col = j
        else:
            raise FormatError(f"{path.name}: unexpected column header {h!r}")
    months = [p for p in cols.values() if p.startswith(str(year))]
    require(months == [f"{year}-{k:02d}" for k in range(1, 13)],
            f"{path.name}: month columns aren't JAN..DEC: {examples(months)}")

    t = Table07(year=year, language=language)
    for r in rows[head_at + 1:]:
        r = list(r) + [None] * (len(header) - len(r))
        code, label = _text(r[0]), _text(r[1])
        if not code and label.upper().startswith(OVERALL_LABELS):
            code = TOTAL
        elif not CODE.match(code):
            continue  # headings, blank rows, footnotes
        weight = _number(r[2])
        series = {p: v for j, p in cols.items() if (v := _number(r[j])) is not None}
        if code in t.labels:
            # The division summary block repeats rows; they must agree with the first copy.
            require(t.indices.get(code) == series and t.weights.get(code) == weight,
                    f"{path.name}: code {code} appears twice with different values")
            continue
        t.labels[code] = label
        if weight is not None:
            t.weights[code] = weight
            t.weight_text[code] = repr(weight)
        if series:
            t.indices[code] = series
        if average_col is not None and (avg := _number(r[average_col])) is not None:
            t.annual_average[code] = avg

    require(TOTAL in t.indices, f"{path.name}: no overall index row")
    require(abs(t.weights.get(TOTAL, 0) - 1000) < 1e-6, f"{path.name}: overall weight isn't 1000")
    divisions = sum(w for c, w in t.weights.items() if len(c) == 2)
    require(abs(divisions - 1000) < 0.5, f"{path.name}: division weights sum to {divisions:.3f}, not 1000")
    values = [v for s in t.indices.values() for v in s.values()]
    require(all(0 < v < 10_000 for v in values), f"{path.name}: index values out of range")
    return t
