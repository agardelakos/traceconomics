"""Read ELSTAT's CPI table 05: the overall index and its twelve-month rate of change, 1959 on.

One file, overwritten every month (`.xlsx`). ELSTAT's own annual inflation rate, as it publishes
it: the site shows it next to the rate we compute from table 07's index levels, and the two are
compared at full precision. The layout, one sheet:

    row 1   title, naming the last year and month ("... during the years 1959-2026 (September)")
    row 2   the base year ("Base Year :2020=100,0")
    then    blocks of eight years side by side, one after the other:
              a row of years, spaced out (" 1 9 5 9 ", "2 0 0 7"), one per pair of columns
              a header row: Month | Index | Rate of change | Index | Rate of change | ...
              a blank row, then months 1 to 12, then " Annual average"
            the last block has empty columns for the years still to come.

Rates are fractions (0.05 is 5%), kept as the file holds them, at full precision. Months after the
latest are empty. Cells outside the blocks are ignored: the file has at least one stray number
between two blocks.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from traceconomics.checks import FormatError, require

# Bump when parsing changes in a way that changes the output: the store then re-parses.
PARSER_VERSION = "1"

TITLE = re.compile(r"twelve-month rates of change of the Overall CPI.*?(\d{4})\s*-\s*(\d{4})\s*\((\w+)\)")
MONTHS = {m: i for i, m in enumerate(("January", "February", "March", "April", "May", "June", "July",
                                      "August", "September", "October", "November", "December"), start=1)}
FIRST_YEAR = 1959


@dataclass
class Table05:
    last: str                                                    # "YYYY-MM", the latest month
    index: dict[str, float] = field(default_factory=dict)        # "YYYY-MM" -> index, 2020=100
    rate: dict[str, float] = field(default_factory=dict)         # "YYYY-MM" -> rate, a fraction
    annual_index: dict[str, float] = field(default_factory=dict)  # "YYYY" -> annual average index
    annual_rate: dict[str, float] = field(default_factory=dict)   # "YYYY" -> its rate, a fraction

    def rows(self):
        """Long form: (freq, unit, ref_period, value). I20 = index 2020=100; RCH_A_FR = annual rate, a fraction."""
        for unit, monthly, annual in (("I20", self.index, self.annual_index), ("RCH_A_FR", self.rate, self.annual_rate)):
            for period, v in monthly.items():
                yield "M", unit, period, v
            for year, v in annual.items():
                yield "A", unit, year, v


def _text(v) -> str:
    return "" if v is None else str(v).strip()


def _number(v) -> float | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    raise FormatError(f"expected a number, got {v!r}")


def _read_rows(path: Path) -> list[list]:
    require(path.suffix.lower() == ".xlsx", f"{path.name}: table 05 is expected as .xlsx")
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        require(len(book.sheetnames) == 1, f"{path.name}: expected one sheet, found {book.sheetnames}")
        return [list(r) for r in book.worksheets[0].iter_rows(values_only=True)]
    finally:
        book.close()


def parse(path: Path) -> Table05:
    rows = _read_rows(path)
    require(len(rows) >= 18, f"{path.name}: only {len(rows)} rows, less than one block")
    title = " ".join(_text(v) for v in rows[0])
    m = TITLE.search(title)
    require(m is not None, f"{path.name}: not table 05's title: {title[:120]!r}")
    first, last_year, last_month = int(m.group(1)), int(m.group(2)), m.group(3).capitalize()
    require(first == FIRST_YEAR, f"{path.name}: the table starts in {first}, not {FIRST_YEAR}")
    require(last_month in MONTHS, f"{path.name}: unknown month in the title: {last_month!r}")
    base = " ".join(_text(v) for v in rows[1])
    require("2020=100" in base, f"{path.name}: base isn't 2020=100 (rebased?): {base!r}")
    t = Table05(last=f"{last_year}-{MONTHS[last_month]:02d}")

    years_seen: list[int] = []
    i = 2
    while i < len(rows):
        row = rows[i]
        labels = [_text(v).replace(" ", "") for v in row]
        if not any(re.fullmatch(r"\d{4}", s) for s in labels):
            i += 1
            continue
        # A block: years over pairs of columns, the header row, a blank row, 12 months, the annual average.
        years = {j: int(s) for j, s in enumerate(labels) if re.fullmatch(r"\d{4}", s)}
        require(all(j % 2 == 1 for j in years), f"{path.name}: row {i + 1}: years not over index columns")
        header = [_text(v) for v in rows[i + 1]]
        require(header[0] == "Month", f"{path.name}: row {i + 2}: expected the Month header, got {header[0]!r}")
        for j in years:
            require(header[j] == "Index" and header[j + 1] == "Rate of change",
                    f"{path.name}: row {i + 2}: columns {j + 1}-{j + 2} aren't Index | Rate of change")
        body = rows[i + 2:i + 16]
        require(len(body) == 14, f"{path.name}: block at row {i + 1} is cut short")
        require(all(v is None for v in body[0]), f"{path.name}: row {i + 3} should be blank")
        for k, r in enumerate(body[1:13], start=1):
            require(_number(r[0]) == k, f"{path.name}: row {i + 3 + k}: expected month {k}, got {r[0]!r}")
        require(_text(body[13][0]) == "Annual average", f"{path.name}: row {i + 16}: no Annual average row")
        for j, year in years.items():
            for k, r in enumerate(body[1:13], start=1):
                period = f"{year}-{k:02d}"
                if (v := _number(r[j])) is not None:
                    t.index[period] = v
                if (v := _number(r[j + 1])) is not None:
                    t.rate[period] = v
            if (v := _number(body[13][j])) is not None:
                t.annual_index[str(year)] = v
            if (v := _number(body[13][j + 1])) is not None:
                t.annual_rate[str(year)] = v
        years_seen += sorted(years.values())
        i += 16

    # The last block runs past the latest year (columns for 2027-2030 in 2026): those stay empty.
    require(bool(years_seen) and years_seen == list(range(FIRST_YEAR, years_seen[-1] + 1)) and years_seen[-1] >= last_year,
            f"{path.name}: years aren't {FIRST_YEAR}-{last_year} or later, in order, one each")
    # Every month up to the latest has an index, none after it; every month from the second year on has a rate.
    expected = [f"{y}-{k:02d}" for y in years_seen for k in range(1, 13) if f"{y}-{k:02d}" <= t.last]
    require(sorted(t.index) == expected, f"{path.name}: index months don't run from {expected[0]} to {t.last}")
    require(sorted(t.rate) == [p for p in expected if p >= f"{FIRST_YEAR + 1}-01"],
            f"{path.name}: rate months don't run from {FIRST_YEAR + 1}-01 to {t.last}")
    require(all(0 < v < 10_000 for v in t.index.values()), f"{path.name}: index values out of range")
    require(all(-1 < v < 10 for v in t.rate.values()), f"{path.name}: rates out of range (not fractions?)")
    return t
