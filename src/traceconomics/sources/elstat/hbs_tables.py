"""Read ELSTAT's household budget survey table 1.03: average monthly spending per household,
by item, for all households and by band of total monthly household income.

The table is the yardstick for our own reading of the survey microdata: rebuilt
from the microdata with the survey weights, the all-households figures must come
out the same.

Two layouts:
- 2008-2023 (.xls): code | label | All households | income bands ...; the rows come in
  blocks by how the goods were acquired, each opened by a total row (code 00, or blank):
  "TOTAL CONSUMPTION", "TOTAL PURCHASES", "FROM OWN PRODUCTION", ...
- 2024 on (.xlsx): YEAR | digits | code | label | way of acquisition (1 = purchases) |
  All households | income bands ...

Only purchases are read: they are what the microdata's "A" columns hold.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import xlrd

from traceconomics.checks import FormatError, require

YEAR_IN_TITLE = re.compile(r"\b(20\d\d)\b")
CODE = re.compile(r"^\d{2,8}$")
ALL = "All households"


@dataclass
class Table103:
    year: int
    groups: list[str]                                             # "All households", then the income bands
    households: dict[str, float] = field(default_factory=dict)    # group -> number of households (weighted)
    purchases: dict[str, dict[str, float]] = field(default_factory=dict)   # code -> group -> euros a month
    labels: dict[str, str] = field(default_factory=dict)
    anomalies: list[str] = field(default_factory=list)            # irregularities in the published table, kept visible


def _rows(path: Path) -> list[list]:
    if path.suffix.lower() == ".xls":
        sheet = xlrd.open_workbook(path, on_demand=True).sheet_by_index(0)
        return [sheet.row_values(i) for i in range(sheet.nrows)]
    if path.suffix.lower() == ".xlsx":
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return [list(r) for r in book[book.sheetnames[0]].iter_rows(values_only=True)]
        finally:
            book.close()
    raise FormatError(f"{path.name}: not an Excel file")


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _number(v, where: str) -> float | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        raise FormatError(f"{where}: expected a number, got {v!r}") from None


def parse(path: Path, year: int) -> Table103:
    """`year`: the survey year, from the file's original name (older titles don't state it)."""
    rows = _rows(path)
    require(len(rows) > 100, f"{path.name}: only {len(rows)} rows")
    title = _text(rows[0][0])
    require("1.03" in title, f"{path.name}: not table 1.03: {title[:80]!r}")
    new_layout = _text(rows[2][0]).upper() == "YEAR"
    if new_layout:
        first_value, code_col, label_col = 5, 2, 3
        stated = _text(rows[4][0])
    else:
        first_value, code_col, label_col = 2, 0, 1
        m = YEAR_IN_TITLE.search(title)
        stated = m.group(1) if m else ""
    require(not stated or int(stated) == year, f"{path.name}: the table says {stated}, expected {year}")

    bands = [_text(v) for v in rows[3][first_value + 1:]]
    groups = [ALL] + [b for b in bands if b]
    require(len(groups) >= 6, f"{path.name}: expected income bands in row 4, found {groups}")
    t = Table103(year=year, groups=groups)

    in_purchases = False
    division = "00"   # the division of the last code seen in the purchases block, to vet codes stored as numbers
    for i, r in enumerate(rows[4:], start=5):
        r = list(r) + [None] * (first_value + len(groups) - len(r))
        code, label = _text(r[code_col]), _text(r[label_col])
        values = r[first_value:first_value + len(groups)]
        if ALL in (label, code) and not CODE.match(code):
            t.households = {g: _number(v, f"{path.name} row {i}") for g, v in zip(groups, values)}
            continue
        if new_layout:
            if not CODE.match(code) or _text(r[4]) != "1":
                continue
        else:
            if code in ("", "00") and label.upper().startswith(("TOTAL", "FROM")):
                code = "00"  # block totals: "00" in some years, blank in others
                in_purchases = label.upper().startswith("TOTAL PURCHASES")
                if not in_purchases:
                    continue
            if not in_purchases or not CODE.match(code):
                continue
            if isinstance(r[code_col], (int, float)):
                # Some years store codes of divisions 10-13 as numbers (e.g. 12312). That is harmless,
                # since they have no leading zero to lose. A code of divisions 01-09 stored as a number
                # would have lost its zero (01111 -> 1111) and could pass for another real code; it
                # shows up as a "10-13" code among codes of a lower division, and is refused.
                ok = code[:2] in ("10", "11", "12", "13") and 0 <= int(code[:2]) - int(division) <= 1
                require(ok, f"{path.name} row {i}: code stored as a number ({r[code_col]!r}) "
                            f"among codes of division {division}; a leading zero may be lost")
            if CODE.match(code):
                division = code[:2]
        if code in t.purchases:
            # A repeated code of up to 5 digits would make the table ambiguous: stop. ELSTAT's
            # 7-digit item codes are national detail with the odd typo (e.g. 2022 repeats 1251001
            # "Life insurance" under 12550 "Other insurance"); keep the first and record it.
            require(len(code) > 5, f"{path.name} row {i}: code {code} twice in the purchases block")
            t.anomalies.append(f"row {i}: code {code} ({label}) repeated; the first occurrence is kept")
            continue
        t.purchases[code] = {g: _number(v, f"{path.name} row {i}") for g, v in zip(groups, values)}
        t.labels[code] = label

    require(t.households.get(ALL), f"{path.name}: no 'All households' count")
    require("00" in t.purchases, f"{path.name}: no total purchases row")
    require(len(t.purchases) > 300, f"{path.name}: only {len(t.purchases)} purchase rows")
    return t
